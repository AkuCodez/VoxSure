"""Build the short submission report as an EDITABLE Word document (same content as the PDF).

    python ml/scripts/build_submission_docx.py --name "Your Name" --regno "23XXX0000" --live https://... --repo https://github.com/...
Output: docs/VoxSure_Project_Report.docx   (text and numbers come from submission_content.py, shared with the PDF)
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from submission_content import L, ROOT, build_blocks  # noqa: E402

NAVY, BLUE, GREY = RGBColor(0x0F, 0x2A, 0x4A), RGBColor(0x16, 0x3A, 0x63), RGBColor(0x44, 0x44, 0x44)
TOKEN = re.compile(r"(<b>|</b>|<i>|</i>)")


def add_runs(par, text, size=None, color=None, bold=False):
    b = i = False
    for part in TOKEN.split(text):
        if part == "<b>":
            b = True
        elif part == "</b>":
            b = False
        elif part == "<i>":
            i = True
        elif part == "</i>":
            i = False
        elif part:
            run = par.add_run(part.replace("&amp;", "&"))
            run.bold, run.italic = bold or b, i
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def page_field(par):
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="[Your Name]")
    ap.add_argument("--regno", default="[Reg. No.]")
    ap.add_argument("--live", default="")
    ap.add_argument("--repo", default="")
    ap.add_argument("--out", default=str(ROOT / "docs" / "VoxSure_Project_Report.docx"))
    a = ap.parse_args()
    links = L(ROOT / "docs" / "links.json") or {}
    live = a.live or links.get("live") or "[live link: add after deployment]"
    repo = a.repo or links.get("repo") or "[GitHub link: add after pushing the code]"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    text_w = 17.0
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    hp = sec.header.paragraphs[0]
    add_runs(hp, f"{a.name} ({a.regno}) | VoxSure Voice-Enabled Chatbot Project Report    Page ", size=8, color=GREY)
    page_field(hp)

    for kind, val in build_blocks(a.name, a.regno, live, repo):
        if kind == "title":
            p = doc.add_paragraph()
            add_runs(p, val, size=17, color=NAVY, bold=True)
            p.paragraph_format.space_after = Pt(2)
        elif kind == "sub":
            p = doc.add_paragraph()
            add_runs(p, val, size=10.5)
            p.paragraph_format.space_after = Pt(1)
        elif kind in ("h1", "h2"):
            p = doc.add_paragraph()
            add_runs(p, val, size=13 if kind == "h1" else 10.5, color=NAVY if kind == "h1" else BLUE, bold=True)
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10 if kind == "h1" else 6), Pt(3)
            p.paragraph_format.keep_with_next = True
        elif kind == "p":
            p = doc.add_paragraph()
            add_runs(p, val, size=10)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(5)
        elif kind in ("note", "caption"):
            p = doc.add_paragraph()
            add_runs(p, val, size=8.5, color=GREY)
            p.paragraph_format.space_after = Pt(4)
        elif kind == "table":
            rows, widths = val["rows"], val["widths"]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style, t.alignment, t.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
            for ri, row in enumerate(rows):
                for ci, txt in enumerate(row):
                    c = t.cell(ri, ci)
                    c.width = Cm(text_w * widths[ci])
                    par = c.paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    add_runs(par, str(txt), size=8.5, bold=(ri == 0))
                    if ri == 0:
                        shade(c, "E3EEFA")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "image":
            doc.add_picture(val["path"], width=Cm(val["width_cm"]))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title, doc.core_properties.author = "VoxSure Project Report", a.name
    doc.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
